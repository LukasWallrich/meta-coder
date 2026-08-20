from io import BytesIO

import pytest
from docx import Document

import meta_coder.manual_drafting as manual_drafting
from meta_coder.manual_drafting import (
    ManualDraftError,
    build_manual_draft_prompt,
    build_manual_draft_schema,
    extract_manual_document_text,
    parse_manual_draft_response,
)


def test_extracts_text_and_tables_from_docx_manual():
    stream = BytesIO()
    document = Document()
    document.add_heading("Affordance Meta-analysis Coding Manual")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Age Mean"
    table.cell(0, 1).text = "Mean age of participants"
    document.save(stream)
    stream.seek(0)

    text = extract_manual_document_text(stream, "manual.docx", max_bytes=1024 * 1024)

    assert "Affordance Meta-analysis Coding Manual" in text
    assert "Age Mean" in text
    assert "Mean age of participants" in text


def test_extracts_page_marked_text_from_pdf(monkeypatch):
    class Page:
        def extract_text(self):
            return "ORIGINAL RESEARCH"

    class Reader:
        is_encrypted = False
        pages = [Page()]

    monkeypatch.setattr(manual_drafting, "PdfReader", lambda _stream: Reader())
    text = extract_manual_document_text(BytesIO(b"%PDF-synthetic"), "manual.pdf", max_bytes=1024)

    assert text.startswith("[Page 1]")
    assert "ORIGINAL RESEARCH" in text


def test_rejects_wrong_document_type_and_spoofed_docx():
    with pytest.raises(ManualDraftError, match="PDF or DOCX"):
        extract_manual_document_text(BytesIO(b"text"), "manual.txt", max_bytes=100)
    with pytest.raises(ManualDraftError, match="not a valid DOCX"):
        extract_manual_document_text(BytesIO(b"not a zip"), "manual.docx", max_bytes=100)


def test_rejects_document_over_upload_limit_before_parsing():
    with pytest.raises(ManualDraftError, match="upload limit"):
        extract_manual_document_text(BytesIO(b"PK" + b"x" * 20), "manual.docx", max_bytes=10)


def test_draft_response_uses_canonical_manual_validation_and_adds_notes():
    manual = parse_manual_draft_response(
        """{
          "name": "attention_review",
          "description": "Attention studies",
          "effect_definition": "Difference between cued and uncued trials",
          "effects": [{
            "name": "cue_type",
            "type": "string",
            "description": "Type of cue",
            "evidence_required": true,
            "levels": [{"value": "valid", "description": "Cue matches target"}]
          }]
        }"""
    )

    assert manual.effect_definition == "Difference between cued and uncued trials"
    assert set(manual.effects) == {"cue_type", "notes"}


def test_invalid_model_draft_is_not_accepted():
    with pytest.raises(ManualDraftError, match="invalid coding-manual draft"):
        parse_manual_draft_response(
            '{"name":"bad","description":"","effect_definition":"","effects":[]}'
        )


def test_schema_dialects_match_provider_requirements():
    gemini = build_manual_draft_schema()
    json_schema = build_manual_draft_schema(dialect="json_schema")

    assert gemini["type"] == "OBJECT"
    assert json_schema["type"] == "object"
    assert json_schema["additionalProperties"] is False
    assert (
        json_schema["properties"]["effects"]["items"]["properties"]["type"]["enum"]
        == ["string", "number", "integer", "boolean"]
    )


def test_prompt_marks_document_and_requires_reviewable_scope():
    prompt = build_manual_draft_prompt("FIELD DEFINITIONS")
    assert prompt.endswith("FIELD DEFINITIONS")
    assert "Do not invent fields" in prompt
    assert "Do not add a notes field" in prompt
